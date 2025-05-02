from flask import Flask, render_template, request
from flask_wtf import FlaskForm
from wtforms import FileField, SubmitField
from wtforms.validators import InputRequired
from werkzeug.utils import secure_filename
import os
from docx import Document

app = Flask(__name__)
app.config['SECRET_KEY'] = 'your_super_secret_key'
app.config['UPLOAD_FOLDER'] = '/Users/andywang/Documents/Projects/report-generator/uploads'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

class UploadForm(FlaskForm):
    file = FileField('Word Document', validators=[InputRequired()])
    submit = SubmitField('Upload')

@app.route('/', methods=['GET', 'POST'])
def upload_file():
    form = UploadForm()
    paragraph_content = []
    table_content = []

    if form.validate_on_submit():
        uploaded_file = form.file.data
        if uploaded_file and uploaded_file.filename.endswith('.docx'):
            filename = secure_filename(uploaded_file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            uploaded_file.save(filepath)

            try:
                document = Document(filepath)

                for paragraph in document.paragraphs:
                    paragraph_content.append(paragraph.text)

                for table in document.tables:
                    table_data = []
                    for row in table.rows:
                        row_data = []
                        for cell in row.cells:
                            row_data.append(cell.text)
                        table_data.append(row_data)
                    table_content.append(table_data)

            except Exception as e:
                paragraph_content = [f"Error reading document: {e}"]
                table_content = []

        elif uploaded_file:
            paragraph_content = ["Please upload a .docx file."]
            table_content = []

        return render_template('index.html', form=form, message="File processed.",
                               paragraph_content='\n'.join(paragraph_content),
                               table_content=table_content)

    return render_template('index.html', form=form, message=None,
                           paragraph_content=None, table_content=None)

if __name__ == '__main__':
    app.run(debug=True)